from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from docx import Document
from docx2pdf import convert

import comtypes.client  # Import comtypes for COM initialization

app = Flask(__name__)
CORS(app)

TEMPLATE_PATH = './templates/agree.docx'  # Update this to match your actual path
OUTPUT_DOCX_PATH = 'tenant_agreement_filled.docx'
OUTPUT_PDF_PATH = 'tenant_agreement_filled.pdf'

def create_tenant_agreement_docx(template_path, output_path, placeholders):
    # Load the document
    doc = Document(template_path)

    # Iterate over paragraphs and replace placeholders
    for paragraph in doc.paragraphs:
        for key, value in placeholders.items():
            if f'{{{key}}}' in paragraph.text:
                paragraph.text = paragraph.text.replace(f'{{{key}}}', str(value))

    # Save the modified document
    doc.save(output_path)

def convert_to_pdf(docx_path, pdf_path):
    # Initialize COM library
    comtypes.CoInitialize()

    # Convert DOCX to PDF
    convert(docx_path, pdf_path)

    # Uninitialize COM library
    comtypes.CoUninitialize()

@app.route('/generate-agreement', methods=['POST'])
def generate_agreement():
    try:
        data = request.json

        placeholders = {
            'agreement_date': data.get('agreement_date', ''),
            'landlord_name': data.get('landlord_name', ''),
            'landlord_address': data.get('landlord_address', ''),
            'tenant_name': data.get('tenant_name', ''),
            'property_address': data.get('property_address', ''),
            'move_in_date': data.get('move_in_date', ''),
            'move_out_date': data.get('move_out_date', ''),
            'rent_amount': data.get('rent_amount', ''),
            'rent_due_day': data.get('rent_due_day', ''),
            'payment_method': data.get('payment_method', ''),
            'security_deposit': data.get('security_deposit', ''),
            'utilities_list': data.get('utilities_list', ''),
            'termination_notice_period': data.get('termination_notice_period', ''),
            'landlord_signature_date': data.get('landlord_signature_date', ''),
            'tenant_signature_date': data.get('tenant_signature_date', '')
        }

        create_tenant_agreement_docx(TEMPLATE_PATH, OUTPUT_DOCX_PATH, placeholders)

        # Convert the DOCX to PDF
        convert_to_pdf(OUTPUT_DOCX_PATH, OUTPUT_PDF_PATH)

        # Send the file to the client
        return send_file(OUTPUT_PDF_PATH, as_attachment=True, download_name='tenant_agreement_filled.pdf')
    except Exception as e:
        return jsonify({"error": f"Error creating tenant agreement: {str(e)}"}), 500

if __name__ == '__main__':
    app.run(debug=True)